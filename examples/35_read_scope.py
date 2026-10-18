"""真实制作小文件验证读取边界；不调用模型或搜索。"""
import sys
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pandas as pd
from docx import Document
from pypdf import PdfWriter
from api.context import set_session_context, set_thread_context, reset_session_context
from api.evidence import EvidenceRegistry, set_evidence, reset_evidence
from tools.upload_file_read_tool import extract_file_details, read_file_content


def main():
    with TemporaryDirectory() as folder:
        root = Path(folder)
        (root / "长文.txt").write_text("开头" + "甲" * 13000 + "末尾未返回", encoding="utf-8")
        text, meta = extract_file_details(root / "长文.txt")
        assert len(text) == 12000 and meta["char_truncated"]
        assert "末尾未返回" not in text and "已截断" in meta["scope"]
        document = Document()
        document.add_paragraph("正文段落")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "表格不读取"
        document.save(root / "说明.docx")
        text, meta = extract_file_details(root / "说明.docx")
        assert text == "正文段落" and meta["tables_omitted"] == 1
        writer = PdfWriter()
        for _ in range(22):
            writer.add_blank_page(width=200, height=200)
        writer.write(root / "页数.pdf")
        _, meta = extract_file_details(root / "页数.pdf")
        assert meta["pages_read"] == 20 and meta["total_pages"] == 22 and meta["empty_pages"] == 20
        with pd.ExcelWriter(root / "订单.xlsx", engine="openpyxl") as book:
            pd.DataFrame({"订单": range(105)}).to_excel(book, sheet_name="订单", index=False)
            pd.DataFrame({"保留": [1]}).to_excel(book, sheet_name="未读表", index=False)
        _, meta = extract_file_details(root / "订单.xlsx")
        assert meta["rows_read"] == 100 and meta["rows_shown"] == 10 and meta["more_rows"]
        assert meta["sheet_count"] == 2 and "总行数未统计" in meta["scope"]
        registry = EvidenceRegistry()
        token = set_evidence(registry)
        path_token = set_session_context(folder)
        thread_token = set_thread_context("scope-example")
        try:
            response = read_file_content.invoke({"filename": "长文.txt"})
            assert "[F1]" in response and "已截断" in response
            assert len(registry.snapshot()) == 1
            read_file_content.invoke({"filename": "长文.txt"})
            assert len(registry.snapshot()) == 1, "相同文件与范围不重复编号"
            assert "读取失败" in read_file_content.invoke({"filename": "../越界.txt"})
        finally:
            reset_session_context(path_token, thread_token)
            reset_evidence(token)
    print("通过：TXT字符截断、Word排除表格、PDF前20页与无OCR、Excel首表100行/展示10行、来源编号与越界拒绝")


if __name__ == "__main__":
    main()
