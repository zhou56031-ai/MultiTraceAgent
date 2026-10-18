from pathlib import Path

import pandas as pd
from docx import Document
from langchain_core.tools import tool
from pypdf import PdfReader

from api.monitor import monitor
from api.evidence import EvidenceRegistry, get_evidence
from utils.path_utils import resolve_session_file


MAX_TEXT_CHARS = 12000


def extract_file_details(path: Path) -> tuple[str, dict]:
    """返回有界正文和读取说明；不把未读取部分说成已经分析。"""
    suffix = path.suffix.lower()
    metadata = {"format": suffix.lstrip("."), "truncated": False}
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
        scope = "UTF-8 文本"
    elif suffix == ".docx":
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        metadata.update(paragraphs=len(document.paragraphs), tables_omitted=len(document.tables))
        scope = f"Word 正文 {len(document.paragraphs)} 段；不读取表格、图片、页眉页脚"
        metadata["truncated"] = bool(document.tables)
    elif suffix == ".pdf":
        reader = PdfReader(path)
        total = len(reader.pages)
        selected = min(total, 20)
        texts = [page.extract_text() or "" for page in reader.pages[:20]]
        text = "\n".join(texts)
        metadata.update(total_pages=total, pages_read=selected, empty_pages=sum(not t.strip() for t in texts))
        metadata["truncated"] = selected < total
        scope = f"PDF 共 {total} 页，仅提取第 1～{selected} 页文字；不做 OCR，图片及扫描页可能无文字"
    elif suffix == ".xlsx":
        with pd.ExcelFile(path, engine="openpyxl") as book:
            sheet = book.sheet_names[0]
            sheet_count = len(book.sheet_names)
            # 第101行仅判断是否还有后续数据，不交给模型；总行数不作猜测。
            sample = pd.read_excel(book, sheet_name=0, nrows=101)
        frame = sample.head(100)
        shown = min(10, len(frame))
        cell_truncated = any(len(str(value)) > 200 for value in frame.head(10).to_numpy().flat)
        text = (
            f"列名：{', '.join(map(str, frame.columns))}\n"
            f"读取行数：{len(frame)}\n"
            f"{frame.head(10).to_string(index=False, max_cols=None, max_colwidth=200)}"
        )
        metadata.update(sheet=sheet, sheet_count=sheet_count, rows_read=len(frame), rows_shown=shown,
                        more_rows=len(sample) > 100, cell_char_limit=200, cell_truncated=cell_truncated)
        metadata["truncated"] = len(frame) > shown or len(sample) > 100 or sheet_count > 1 or cell_truncated
        scope = (f"Excel 共 {sheet_count} 个工作表，仅首表「{sheet}」；最多取 100 行数据，"
                 f"本次取 {len(frame)} 行、正文展示前 {shown} 行；总行数未统计；单元格展示最多 200 字符")
    else:
        raise ValueError(f"不支持的文件类型：{suffix or '无扩展名'}")

    body = text[:MAX_TEXT_CHARS]
    char_truncated = len(text) > len(body)
    metadata.update(extracted_chars=len(text), returned_chars=len(body), char_truncated=char_truncated)
    metadata["truncated"] = metadata["truncated"] or char_truncated
    scope += f"；所选范围提取 {len(text)} 字符，返回前 {len(body)} 字符"
    if char_truncated:
        scope += "（正文已截断，后续内容未返回）"
    metadata["scope"] = scope
    return body, metadata


def extract_file_text(path: Path) -> str:
    """兼容前面教程的普通函数入口；返回有界的提取文字。"""
    text, _ = extract_file_details(path)
    return text


@tool
def read_file_content(filename: str) -> str:
    """读取当前会话中的 txt、md、docx、pdf 或 xlsx 文件。"""
    try:
        target = resolve_session_file(filename)
        if not target.is_file():
            return f"错误：文件不存在：{target.name}"
        monitor.report_tool("read_file_content", {"filename": target.name})
        text, metadata = extract_file_details(target)
        registry = get_evidence() or EvidenceRegistry()
        source = registry.add("file", (filename, target.stat().st_mtime_ns, metadata["scope"]),
                              {"filename": filename, **metadata})
        monitor.report_sources(registry.snapshot())
        return (f"[{source['id']}] 文件：{filename}\n【读取范围】{metadata['scope']}\n"
                "回答时引用此文件编号；仅据已返回内容作答，不声称已通读未读取部分。\n\n"
                + (text if text.strip() else "文件中没有可提取的文字；扫描 PDF 需要另行 OCR"))
    except (ValueError, RuntimeError, OSError, UnicodeError) as exc:
        return f"读取失败：{exc}"
    except Exception as exc:
        return f"读取失败：{type(exc).__name__}"


